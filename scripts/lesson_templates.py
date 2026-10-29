"""Profile validation and OOXML filling; untouched ZIP members remain unchanged."""
from copy import deepcopy
import io
import json
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from lxml import etree as E

ROOT = Path(__file__).resolve().parents[1]
PROFILES = json.loads((ROOT / 'templates/profiles.json').read_text(encoding='utf-8'))
NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}


def tag(name):
    return '{' + NS['w'] + '}' + name


def el(name, **attrs):
    node = E.Element(tag(name))
    for key, value in attrs.items():
        node.set(tag(key), str(value))
    return node


def text_required(value, label):
    if not isinstance(value, str) or not value.strip():
        raise ValueError(f'{label} 必須是非空字串')


def positive_int(value, label):
    if type(value) is not int or value < 1:
        raise ValueError(f'{label} 必須是正整數')


def validate_payload(data, profile):
    for key in ('metadata', 'fields') + (('preparation',) if profile == 'digital-ai' else ()):
        if not isinstance(data.get(key), dict):
            raise ValueError(f'{key} 必須是物件')
    arrays = ('periods', 'rubrics', 'reflections', 'appendices', 'sources')
    if profile == 'digital-ai':
        arrays += ('checklist', 'integration')
    for key in arrays:
        if key in data and not isinstance(data[key], list):
            raise ValueError(f'{key} 必須是陣列')
    if type(data.get('schema_version')) is not int or data['schema_version'] != 1:
        raise ValueError('schema_version 必須為 1')
    if data.get('template_id') != profile:
        raise ValueError('JSON template_id 與 --template 不一致')
    meta = data['metadata']
    for key in ('topic', 'subject', 'grade', 'designer'):
        text_required(meta.get(key), 'metadata.' + key)
    for key in ('period_count', 'minutes_per_period'):
        positive_int(meta.get(key), 'metadata.' + key)
    for key in PROFILES[profile]['required_fields']:
        text_required(data['fields'].get(key), 'fields.' + key)
    periods = data['periods']
    if not isinstance(periods, list) or len(periods) != meta['period_count']:
        raise ValueError('periods 數量必須等於 period_count')
    for number, period in enumerate(periods, 1):
        if not isinstance(period, dict):
            raise ValueError(f'第{number}節必須是物件')
        for key in ('title', 'competency_link'):
            text_required(period.get(key), f'第{number}節 {key}')
        stages = period['stages']
        if not isinstance(stages, list) or not stages:
            raise ValueError('每節必須有活動階段')
        for stage in stages:
            if not isinstance(stage, dict):
                raise ValueError(f'第{number}節 stage 必須是物件')
            for key in ('name', 'activity', 'assessment', 'strategy'):
                text_required(stage.get(key), f'第{number}節 {key}')
            positive_int(stage.get('minutes'), f'第{number}節分鐘數')
            if not isinstance(stage.get('notes', ''), str):
                raise ValueError('notes 必須是字串')
        if sum(s['minutes'] for s in stages) != meta['minutes_per_period']:
            raise ValueError(f'第{number}節階段時間總和不等於每節分鐘數')
    if data.get('reflection_status') not in ('planned', 'observed'):
        raise ValueError('reflection_status 必須是 planned 或 observed')
    if not isinstance(data['reflections'], list) or not data['reflections']:
        raise ValueError('reflections 至少一段')
    for text in data['reflections']:
        text_required(text, 'reflection')
    rubrics = data['rubrics']
    if not isinstance(rubrics, list) or not rubrics:
        raise ValueError('rubrics 至少一個評量表')
    for rubric in rubrics:
        if not isinstance(rubric, dict):
            raise ValueError('rubric 必須是物件')
        text_required(rubric.get('title'), 'rubric.title')
        headers, rows = rubric['headers'], rubric['rows']
        if not isinstance(headers, list) or not headers or not isinstance(rows, list) or not rows:
            raise ValueError('評量表需要表頭及資料列')
        for row in [headers] + rows:
            if not isinstance(row, list) or len(row) != len(headers):
                raise ValueError('評量表每列欄位數必須一致')
            for text in row:
                text_required(text, 'rubric cell')
    for item in data.get('appendices', []):
        if not isinstance(item, dict):
            raise ValueError('appendix 必須是物件')
        text_required(item.get('title'), 'appendix.title')
        text_required(item.get('text'), 'appendix.text')
    for item in data.get('sources', []):
        if not isinstance(item, dict):
            raise ValueError('source 必須是物件')
        text_required(item.get('title'), 'source.title')
        text_required(item.get('used_for'), 'source.used_for')
        if not isinstance(item.get('url', ''), str):
            raise ValueError('source.url 必須是字串')
    if profile == 'cross-domain':
        if meta.get('course_type') not in ('議題融入式課程', '議題主題式課程'):
            raise ValueError('course_type 不在跨域範本選項中')
        if meta.get('implementation') not in ('領域學習課程', '彈性學習課程'):
            raise ValueError('implementation 不在跨域範本選項中')
    else:
        if len(rubrics) != 1 or len(rubrics[0]['headers']) != 4 or len(rubrics[0]['rows']) != 3:
            raise ValueError('digital-ai 評量表需要四欄、三項規準加一列表頭')
        if len(data['reflections']) != 3:
            raise ValueError('digital-ai 需三段省思')
        for key in ('teacher', 'student', 'teacher_evidence', 'student_evidence'):
            text_required(data['preparation'].get(key), 'preparation.' + key)
        if not data.get('checklist') or not data.get('integration'):
            raise ValueError('digital-ai 需 checklist 與 integration')
        for item in data['checklist']:
            if not isinstance(item, dict):
                raise ValueError('checklist 項目必須是物件')
            text_required(item.get('item'), 'checklist.item')
            if type(item.get('checked')) is not bool or item.get('light') not in ('', '綠燈', '黃燈', '紅燈'):
                raise ValueError('checked 需布林值，light 使用空字串或綠燈／黃燈／紅燈')
        for item in data['integration']:
            if not isinstance(item, dict):
                raise ValueError('integration 項目必須是物件')
            text_required(item.get('item'), 'integration.item')
            if type(item.get('ai')) is not bool or not isinstance(item.get('dimensions'), dict):
                raise ValueError('integration 需 ai 布林值及 dimensions 物件')
            for key, text in item['dimensions'].items():
                if key not in tuple('ABCDEFG'):
                    raise ValueError('融入向度只接受 A 至 G')
                text_required(text, 'integration.dimensions.' + key)


def cell(table, row_number, column):
    row = table.findall('w:tr', NS)[row_number]
    offset = 0
    for tc in row.findall('w:tc', NS):
        span = tc.find('w:tcPr/w:gridSpan', NS)
        size = int(span.get(tag('val'))) if span is not None else 1
        if offset <= column < offset + size:
            return tc
        offset += size
    raise ValueError(f'找不到儲存格 {row_number}, {column}')


def node_text(node):
    return ''.join(node.xpath('.//w:t/text()', namespaces=NS))


def write_text(node, text, size=None, font=None):
    is_cell = node.tag == tag('tc')
    p = node.find('w:p', NS) if is_cell else node
    pp = deepcopy(p.find('w:pPr', NS)) if p is not None else None
    rp = deepcopy(p.find('.//w:rPr', NS)) if p is not None else None
    for child in list(node):
        if not is_cell or child.tag != tag('tcPr'):
            node.remove(child)
    p = el('p') if is_cell else node
    if pp is not None:
        p.append(pp)
    r = el('r'); rp = rp if rp is not None else el('rPr')
    if font:
        fonts = rp.find('w:rFonts', NS)
        if fonts is None:
            fonts = el('rFonts'); rp.append(fonts)
        for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(tag(attr), font)
    if size:
        for name in ('sz', 'szCs'):
            entry = rp.find('w:' + name, NS)
            if entry is None:
                entry = el(name); rp.append(entry)
            entry.set(tag('val'), str(size * 2))
    r.append(rp)
    for i, line in enumerate(text.split('\n')):
        if i:
            r.append(el('br'))
        t = el('t'); t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve'); t.text = line; r.append(t)
    p.append(r)
    if is_cell:
        node.append(p)


def fill(table, row, col, text, profile):
    write_text(cell(table, row, col), text, 12 if profile == 'digital-ai' else 14,
               None if profile == 'digital-ai' else '標楷體')


def activity_rows(table, start, old_count, new_count):
    rows = table.findall('w:tr', NS)
    prototype = deepcopy(rows[start]); position = list(table).index(rows[start])
    for row in rows[start:start + old_count]:
        table.remove(row)
    for i in range(new_count):
        table.insert(position + i, deepcopy(prototype))


def heading(body, text):
    for p in body.findall('w:p', NS):
        if node_text(p).strip() == text:
            return p
    raise ValueError(f'範本缺少標題：{text}')


def after(anchor, text, profile):
    p = el('p'); write_text(p, text, 12 if profile == 'digital-ai' else 14,
                            None if profile == 'digital-ai' else '標楷體')
    anchor.addnext(p)
    return p


def add_rubric(rubric, profile):
    table = el('tbl'); props = el('tblPr'); table.append(props)
    width = 14600 if profile == 'digital-ai' else 10200
    props.append(el('tblW', w=width, type='dxa'))
    borders = el('tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        borders.append(el(side, val='single', sz=4, color='808080'))
    props.append(borders); grid = el('tblGrid'); table.append(grid)
    col_width = width // len(rubric['headers'])
    for _ in rubric['headers']:
        grid.append(el('gridCol', w=col_width))
    for i, values in enumerate([rubric['headers']] + rubric['rows']):
        row = el('tr'); table.append(row)
        if i == 0:
            rp = el('trPr'); rp.append(el('tblHeader')); row.append(rp)
        for value in values:
            tc = el('tc'); row.append(tc); tcp = el('tcPr')
            tcp.append(el('tcW', w=col_width, type='dxa')); tc.append(tcp)
            if i == 0:
                tcp.append(el('shd', fill='E7E6E6', val='clear'))
            write_text(tc, value, 12 if profile == 'digital-ai' else 14,
                       None if profile == 'digital-ai' else '標楷體')
    return table


def reflections(data):
    prefix = '課後待填／預期觀察' if data['reflection_status'] == 'planned' else '依提供之實際教學紀錄'
    return [prefix + '：' + t for t in data['reflections']]


def fill_cross(body, table, data):
    profile = 'cross-domain'; m = data['metadata']; f = data['fields']
    for p in body.findall('w:p', NS):
        text = node_text(p).strip()
        if text.startswith('主題名稱：'):
            write_text(p, '主題名稱：' + m['topic'])
        elif text.startswith('設計者：'):
            write_text(p, '設計者：' + m['designer'])
    for title, key in [('設計動機', 'motivation'), ('期望之學生學習歷程', 'learning_journey'),
                       ('呼應核心素養的說明', 'competency_explanation'), ('主題架構說明', 'theme_structure')]:
        after(heading(body, title), f[key], profile)
    def choices(selected, options):
        return '\n'.join(('■' if v == selected else '□') + v for v in options)
    values = {(0, 3): m['topic'], (1, 3): choices(m['course_type'], ('議題融入式課程', '議題主題式課程')),
              (1, 6): m['designer'], (2, 3): choices(m['implementation'], ('領域學習課程', '彈性學習課程')),
              (2, 6): m['subject'], (3, 3): m['grade'],
              (3, 6): f'{m["period_count"]}節，每節{m["minutes_per_period"]}分鐘'}
    slots = {(5, 3): 'subject_prior_knowledge', (6, 3): 'digital_prior_knowledge',
             (7, 3): 'competency_dimensions', (7, 6): 'core_competencies',
             (8, 3): 'domain_competencies', (8, 6): 'issue_competencies',
             (9, 3): 'learning_performance', (9, 6): 'issue_themes',
             (10, 3): 'learning_content', (10, 6): 'issue_content',
             (11, 3): 'objectives', (12, 3): 'assessment_strategy', (13, 3): 'resources'}
    values.update({pos: f[key] for pos, key in slots.items()})
    for (row, col), text in values.items():
        fill(table, row, col, text, profile)
    activity_rows(table, 16, 6, len(data['periods']))
    for i, period in enumerate(data['periods']):
        stages = period['stages']; row = 16 + i
        text = f'第{i + 1}節　{period["title"]}\n' + '\n\n'.join(s['name'] + '\n' + s['activity'] for s in stages)
        fill(table, row, 1, text + '\n\n核心素養連結：' + period['competency_link'], profile)
        for col, key in ((4, 'minutes'), (5, 'assessment'), (7, 'strategy'), (8, 'notes')):
            text = '\n\n'.join(str(s[key]) + '分' if key == 'minutes' else
                               (s.get(key, '') if key == 'strategy' else s['name'] + '：' + s.get(key, ''))
                               for s in stages)
            fill(table, row, col, text, profile)
    fill(table, 17 + len(data['periods']), 1, '\n\n'.join(reflections(data)), profile)
    anchor = heading(body, '附錄')
    for rubric in data['rubrics']:
        anchor = after(anchor, rubric['title'], profile)
        new = add_rubric(rubric, profile); anchor.addnext(new); anchor = new
    return anchor


def item_row(table, column, item, start=2):
    matches = [i for i in range(start, len(table.findall('w:tr', NS))) if node_text(cell(table, i, column)).strip() == item]
    if len(matches) != 1:
        raise ValueError(f'找不到唯一的範本項目：{item}')
    return matches[0]


def fill_digital(body, tables, data):
    profile = 'digital-ai'; m = data['metadata']; f = data['fields']; t0, flow, checklist, matrix = tables
    values = {(0, 2): m['subject'], (0, 7): m['designer'], (2, 2): m['grade'],
              (2, 3): m.get('class_name', '待填'), (2, 4): m.get('student_count', '待填'),
              (2, 7): f'{m["period_count"]}節，每節{m["minutes_per_period"]}分鐘', (3, 2): m['topic'],
              (4, 2): f['rationale'], (7, 2): f['core_competencies'], (7, 5): f['digital_ai_competencies'],
              (10, 2): f['learning_performance'], (11, 2): f['learning_content'], (12, 2): f['issue_integration']}
    for (row, col), text in values.items():
        fill(t0, row, col, str(text), profile)
    for row, key in enumerate(('connections', 'objectives', 'tasks', 'tools', 'equipment', 'digital_prior_knowledge', 'resources')):
        fill(flow, row, 2, f[key], profile)
    for row, key in ((10, 'teacher'), (11, 'student')):
        fill(flow, row, 1, data['preparation'][key], profile)
        fill(flow, row, 3, data['preparation'][key + '_evidence'], profile)
    activity_rows(flow, 12, 3, len(data['periods']))
    for i, period in enumerate(data['periods']):
        fill(flow, 12 + i, 0, f'第{i + 1}節', profile)
        text = period['title'] + '\n' + '\n\n'.join(f'{s["name"]}（{s["minutes"]}分鐘）\n{s["activity"]}\n教學策略：{s["strategy"]}' for s in period['stages'])
        fill(flow, 12 + i, 1, text + '\n核心素養連結：' + period['competency_link'], profile)
        fill(flow, 12 + i, 3, '\n\n'.join(s['name'] + '：' + s['assessment'] + '\n' + s.get('notes', '') for s in period['stages']), profile)
    for row in range(2, 19):
        fill(checklist, row, 2, '', profile); fill(checklist, row, 3, '', profile)
    seen = set()
    for item in data['checklist']:
        row = item_row(checklist, 0, item['item'])
        if row in seen:
            raise ValueError('checklist 項目重複')
        seen.add(row); fill(checklist, row, 2, 'V' if item['checked'] else '', profile)
        fill(checklist, row, 3, item['light'], profile)
    fill(matrix, 0, 3, '教學方案名稱：' + m['topic'], profile)
    seen = set()
    for item in data['integration']:
        row = item_row(matrix, 1, item['item'])
        if row in seen:
            raise ValueError('integration 項目重複')
        seen.add(row); fill(matrix, row, 2, 'V' if item['ai'] else '', profile)
        for key, text in item['dimensions'].items():
            fill(matrix, row, 3 + ord(key) - ord('A'), text, profile)
    heading(body, '學習評量規準').addnext(add_rubric(data['rubrics'][0], profile))
    anchor = heading(body, '教師教學反思')
    for text in reflections(data):
        anchor = after(anchor, text, profile)
    for rp in body.findall('.//w:rPr', NS):
        for name in ('sz', 'szCs'):
            size = rp.find('w:' + name, NS)
            if size is None:
                size = el(name); rp.append(size)
            size.set(tag('val'), '24')
    return anchor


def convert_json(data, output, requested=None, source=None, check_only=False):
    if not isinstance(data, dict):
        raise ValueError('JSON 根節點必須為物件')
    profile = requested or data.get('template_id')
    if profile not in PROFILES:
        raise ValueError('未知 template_id，請用 --list-templates 查詢')
    validate_payload(data, profile)
    output = Path(output).resolve(); template = ROOT / PROFILES[profile]['file']
    protected = {p.resolve() for folder in ('templates', 'examples', 'references') for p in (ROOT / folder).rglob('*') if p.is_file()}
    if output in protected or (source and output == Path(source).resolve()):
        raise ValueError('不得覆寫範本、參考檔或來源')
    if output.suffix.lower() != '.docx':
        raise ValueError('輸出副檔名必須為 .docx')
    if not check_only and output.exists():
        raise ValueError('輸出檔已存在，請使用新的檔名')
    with ZipFile(template) as original:
        document = E.fromstring(original.read('word/document.xml'))
        body = document.find('w:body', NS); tables = body.findall('w:tbl', NS)
        shapes = [(len(t.findall('w:tr', NS)), len(t.findall('w:tblGrid/w:gridCol', NS))) for t in tables]
        if shapes != [tuple(s) for s in PROFILES[profile]['source_shapes']]:
            raise ValueError('範本結構已變更，需要重新建立欄位對應')
        anchor = fill_cross(body, tables[0], data) if profile == 'cross-domain' else fill_digital(body, tables, data)
        for item in data.get('appendices', []):
            anchor = after(anchor, item['title'] + '\n' + item['text'], profile)
        if data.get('sources'):
            after(anchor, '資料來源\n' + '\n'.join(s['title'] + '　' + s.get('url', '') + '\n用途：' + s['used_for'] for s in data['sources']), profile)
        xml = E.tostring(document, xml_declaration=True, encoding='UTF-8', standalone=True)
        stream = io.BytesIO()
        with ZipFile(stream, 'w') as result:
            for entry in original.infolist():
                result.writestr(entry, xml if entry.filename == 'word/document.xml' else original.read(entry.filename))
    result_bytes = stream.getvalue(); verify_output(result_bytes, data, profile)
    if not check_only:
        output.parent.mkdir(parents=True, exist_ok=True)
        with output.open('xb') as target:
            target.write(result_bytes)


def verify_output(result_bytes, data, profile):
    doc = Document(io.BytesIO(result_bytes)); count = data['metadata']['period_count']
    shapes = [(len(t.rows), len(t.columns)) for t in doc.tables]
    expected = ([(13, 8), (12 + count, 4), (4, 4), (19, 4), (15, 10)] if profile == 'digital-ai' else
                [(18 + count, 9)] + [(1 + len(r['rows']), len(r['headers'])) for r in data['rubrics']])
    if shapes != expected:
        raise ValueError(f'輸出表格結構錯位：{shapes}，預期 {expected}')
    for section in doc.sections:
        if (section.page_width > section.page_height) != (profile == 'digital-ai'):
            raise ValueError('輸出頁面方向不符範本')
    flow = doc.tables[1] if profile == 'digital-ai' else doc.tables[0]
    start, column = (12, 1) if profile == 'digital-ai' else (16, 1)
    for i, period in enumerate(data['periods']):
        text = flow.cell(start + i, column).text
        for stage in period['stages']:
            if stage['activity'] not in text:
                raise ValueError(f'第{i + 1}節內容遺漏')
