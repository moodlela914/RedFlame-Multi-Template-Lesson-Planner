import copy
import hashlib
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
CLI = ROOT / 'scripts' / 'markdown_to_docx.py'


def lesson(profile='cross-domain', count=7):
    fields = {key: f'測試內容 {key}' for key in (
        'motivation learning_journey competency_explanation theme_structure rationale '
        'subject_prior_knowledge digital_prior_knowledge digital_ai_competencies competency_dimensions '
        'core_competencies domain_competencies learning_performance learning_content '
        'issue_competencies issue_themes issue_content objectives assessment_strategy '
        'resources connections tasks tools equipment issue_integration'
    ).split()}
    return {
        'schema_version': 1, 'template_id': profile,
        'metadata': {'topic': '測試課程資料判讀', 'subject': '資訊與自然',
                     'grade': '六年級', 'designer': '待填', 'period_count': count,
                     'minutes_per_period': 40, 'course_type': '議題融入式課程',
                     'implementation': '領域學習課程'},
        'fields': fields,
        'preparation': {'teacher': '教師準備測試教材', 'student': '學生準備測試資料',
                        'teacher_evidence': '教材清單', 'student_evidence': '課前紀錄'},
        'periods': [{'title': f'單元活動{i}', 'competency_link': f'素養連結{i}',
                     'stages': [{'name': '準備活動', 'minutes': 5, 'activity': f'導入{i}',
                                 'assessment': '口頭問答', 'strategy': '提問', 'notes': ''},
                                {'name': '發展活動', 'minutes': 25, 'activity': f'關鍵任務{i}',
                                 'assessment': f'學習證據{i}', 'strategy': '合作探究', 'notes': '資料來源核對'},
                                {'name': '綜合活動', 'minutes': 10, 'activity': f'統整{i}',
                                 'assessment': '離堂單', 'strategy': '發表', 'notes': ''}]}
                    for i in range(1, count + 1)],
        'rubrics': [{'title': '任務評量規準', 'headers': ['評量項目', '優秀', '良好', '待加強'],
                     'rows': [['知識', '能正確解釋', '部分解釋', '需要協助'],
                              ['技能', '能獨立完成', '提示後完成', '需要示範'],
                              ['態度', '主動合作', '參與合作', '需要提醒']]}],
        'reflection_status': 'planned',
        'reflections': ['課後待填：觀察理解情形', '課後待填：觀察操作情形', '課後待填：檢視教學時間'],
        'checklist': [{'item': '教學目標', 'checked': True, 'light': '黃燈'}],
        'integration': [{'item': '1.共同備課', 'ai': True, 'dimensions': {'A': '整理可查證的引導問題'}}],
        'appendices': [{'title': '測試學習單', 'text': '請記錄一項可查證的資料。'}],
        'sources': [{'title': '使用者提供教材', 'url': '', 'used_for': '課程設計'}],
    }


class TemplateTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.folder = Path(self.tmp.name)

    def run_cli(self, payload, *args):
        source = self.folder / 'lesson.json'
        source.write_text(json.dumps(payload, ensure_ascii=False), encoding='utf-8')
        output = self.folder / 'output.docx'
        result = subprocess.run([sys.executable, str(CLI), str(source), str(output), *args],
                                capture_output=True, text=True, encoding='utf-8')
        return result, output

    def test_lists_both_templates(self):
        r = subprocess.run([sys.executable, str(CLI), '--list-templates'], capture_output=True, text=True, encoding='utf-8')
        self.assertEqual(r.returncode, 0, r.stderr + r.stdout)
        self.assertIn('digital-ai', r.stdout)
        self.assertIn('cross-domain', r.stdout)

    def test_cross_domain_seven_periods_and_source_fidelity(self):
        source = ROOT / 'templates' / '資訊議題跨域教案(空白).docx'
        source_hash = hashlib.sha256(source.read_bytes()).hexdigest() if source.exists() else ''
        r, output = self.run_cli(lesson(), '--template', 'cross-domain')
        self.assertEqual(r.returncode, 0, r.stderr + r.stdout)
        d = Document(output)
        self.assertLess(d.sections[0].page_width, d.sections[0].page_height)
        self.assertEqual((len(d.tables[0].rows), len(d.tables[0].columns)), (25, 9))
        text = '\n'.join(p.text for p in d.paragraphs) + '\n' + '\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
        for i in range(1, 8):
            self.assertIn(f'關鍵任務{i}', text)
            self.assertIn(f'學習證據{i}', text)
        self.assertNotIn('E起看星空', text)
        self.assertIn('設計者：待填', text)
        self.assertNotIn('☑已經會議', text)
        self.assertEqual(hashlib.sha256(source.read_bytes()).hexdigest(), source_hash)
        with ZipFile(source) as a, ZipFile(output) as b:
            self.assertEqual(set(a.namelist()), set(b.namelist()))
            for name in a.namelist():
                if name != 'word/document.xml':
                    self.assertEqual(a.read(name), b.read(name), name)
        self.assertIn('標楷體', d.tables[0].cell(0, 3)._tc.xml)
        self.assertIn('w:val="28"', d.tables[0].cell(0, 3)._tc.xml)

    def test_digital_seven_periods_retains_required_tables(self):
        r, output = self.run_cli(lesson('digital-ai'), '--template', 'digital-ai')
        self.assertEqual(r.returncode, 0, r.stderr + r.stdout)
        d = Document(output)
        self.assertTrue(all(s.page_width > s.page_height for s in d.sections))
        self.assertEqual([(len(t.rows), len(t.columns)) for t in d.tables],
                         [(13, 8), (19, 4), (4, 4), (19, 4), (15, 10)])
        self.assertIn('關鍵任務7', d.tables[1].cell(18, 1).text)
        self.assertIn('學習證據7', d.tables[1].cell(18, 3).text)
        self.assertIn('w:val="24"', d.tables[1].cell(18, 1)._tc.xml)

    def test_invalid_inputs_leave_no_output(self):
        bad = []
        p = lesson(); del p['fields']['objectives']; bad.append(p)
        p = lesson(); p['periods'].pop(); bad.append(p)
        p = lesson(); p['periods'][0]['stages'][0]['minutes'] = 9; bad.append(p)
        p = lesson(); p['metadata']['period_count'] = True; bad.append(p)
        p = lesson(); p['periods'][0]['stages'][0]['activity'] = ''; bad.append(p)
        p = lesson(); p['template_id'] = 'unknown'; bad.append(p)
        p = lesson('digital-ai'); p['checklist'][0]['item'] = '不存在的項目'; bad.append(p)
        p = lesson('digital-ai'); p['integration'][0]['dimensions'] = {'Z': '不支援'}; bad.append(p)
        for index, payload in enumerate(bad):
            with self.subTest(case=index):
                r, output = self.run_cli(payload)
                self.assertNotEqual(r.returncode, 0)
                self.assertFalse(output.exists())

    def test_conflicting_profile_rejected(self):
        r, output = self.run_cli(lesson(), '--template', 'digital-ai')
        self.assertNotEqual(r.returncode, 0)
        self.assertFalse(output.exists())

    def test_check_only_validates_without_creating_file(self):
        for profile in ('digital-ai', 'cross-domain'):
            with self.subTest(profile=profile):
                r, output = self.run_cli(lesson(profile), '--check-only')
                self.assertEqual(r.returncode, 0, r.stderr + r.stdout)
                self.assertFalse(output.exists())

    def test_one_period_has_no_blank_extra_activity_rows(self):
        for profile, expected in [('cross-domain', (19, 9)), ('digital-ai', (13, 4))]:
            with self.subTest(profile=profile):
                r, output = self.run_cli(lesson(profile, 1))
                self.assertEqual(r.returncode, 0, r.stderr + r.stdout)
                d = Document(output)
                table = d.tables[0 if profile == 'cross-domain' else 1]
                self.assertEqual((len(table.rows), len(table.columns)), expected)
                output.unlink()

    def test_markdown_with_fourth_period_is_rejected(self):
        source = self.folder / 'old.md'
        source.write_text('# 第四節\n|教學重點|活動|證據|\n|---|---|---|\n|測試|測試|測試|', encoding='utf-8')
        output = self.folder / 'old.docx'
        r = subprocess.run([sys.executable, str(CLI), str(source), str(output)],
                           capture_output=True, text=True, encoding='utf-8')
        self.assertNotEqual(r.returncode, 0)
        self.assertFalse(output.exists())

    def test_malformed_containers_report_errors_without_tracebacks(self):
        bad = []
        p = lesson(); p['metadata'] = []; bad.append(p)
        p = lesson(); p['periods'][0] = None; bad.append(p)
        p = lesson(); p['periods'][0]['stages'][0] = '文字'; bad.append(p)
        p = lesson(); p['rubrics'][0] = []; bad.append(p)
        p = lesson(); p['appendices'] = {}; bad.append(p)
        for index, payload in enumerate(bad):
            with self.subTest(case=index):
                r, output = self.run_cli(payload)
                self.assertNotEqual(r.returncode, 0)
                self.assertNotIn('Traceback', r.stderr)
                self.assertFalse(output.exists())

    def test_existing_output_is_not_overwritten(self):
        output = self.folder / 'output.docx'
        output.write_bytes(b'existing user document')
        r, _ = self.run_cli(lesson())
        self.assertNotEqual(r.returncode, 0)
        self.assertEqual(output.read_bytes(), b'existing user document')

    def test_existing_markdown_examples_still_convert(self):
        for source in ROOT.glob('examples/*/*.md'):
            with self.subTest(source=source):
                out = self.folder / (source.parent.name + '.docx')
                r = subprocess.run([sys.executable, str(CLI), str(source), str(out)],
                                   capture_output=True, text=True, encoding='utf-8')
                self.assertEqual(r.returncode, 0, r.stderr + r.stdout)
                self.assertEqual(len(Document(out).tables), 5)


if __name__ == '__main__':
    unittest.main()
